"""Extract the verified PRJNA260535 skin RNA-seq workbook for beta10 genes.

Requires openpyxl. Run from the canonical repository root. Original source is
archived under data/reference/external_t007/ with URL/SHA in source_manifest.tsv.
No external samples are added to the primary 54-sample GSE98923 matrix.
"""

from __future__ import annotations

import csv
import hashlib
import re
from collections import Counter
from pathlib import Path

import openpyxl


ROOT = Path.cwd()
if not (ROOT / "AGENTS.md").is_file():
    raise RuntimeError("Run from the canonical repository root")

reference = ROOT / "data" / "reference" / "external_t007"
manifest = list(csv.DictReader((reference / "source_manifest.tsv").open(encoding="utf-8"), delimiter="\t"))
if len(manifest) != 3 or len({entry["File"] for entry in manifest}) != 3:
    raise ValueError("Unexpected T-007 source manifest")
for entry in manifest:
    path = reference / entry["File"]
    if not path.is_file() or path.stat().st_size != int(entry["Bytes"]):
        raise ValueError(f"Missing or changed original source: {path}")
    with path.open("rb") as handle:
        sha = hashlib.file_digest(handle, "sha256").hexdigest()
    if sha != entry["SHA256"]:
        raise ValueError(f"Source SHA-256 mismatch: {path}")

selected_path = ROOT / "results" / "beta10" / "tables" / "selected_genes.txt"
selected = selected_path.read_text(encoding="utf-8").split()
if len(selected) != 3050 or len(set(selected)) != 3050:
    raise ValueError("Unexpected beta10 gene universe")
selected_set = set(selected)

source = reference / "PRJNA260535_additional_file_2_log2CPM.xlsx"
workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
if workbook.sheetnames != ["Additional file 9 -log2cpm.csv"]:
    raise ValueError("Unexpected RNA-seq worksheet")
sheet = workbook.active
if sheet.max_row != 16607 or sheet.max_column != 86:
    raise ValueError("Unexpected RNA-seq worksheet dimensions")
rows = sheet.iter_rows(values_only=True)
header = next(rows)
if header[:2] != (None, "V1"):
    raise ValueError("Unexpected RNA-seq header")
samples = header[2:]
if len(samples) != 84 or len(set(samples)) != 84:
    raise ValueError("Expected 84 unique sample names")

names = {
    "CF": "Cabernet Franc",
    "CSW": "Cabernet Sauvignon",
    "MEW": "Merlot",
    "PNW": "Pinot noir",
    "CDW": "Chardonnay",
    "SB": "Sauvignon Blanc",
    "SMW": "Semillon",
}
metadata = []
for name in samples:
    match = re.fullmatch(r"(CF|CSW|MEW|PNW|CDW|SB|SMW)([123])\.(20|22|24|26)", name)
    if match is None:
        raise ValueError(f"Unexpected RNA-seq sample name: {name}")
    prefix, replicate, brix = match.groups()
    metadata.append((name, names[prefix], int(brix), int(replicate), 2012, "berry skin"))
if set(Counter((row[1], row[2]) for row in metadata).values()) != {3}:
    raise ValueError("RNA-seq cultivar/Brix replicate design is not balanced")

genes_seen: set[str] = set()
selected_rows = []
missing_values = 0
for row in rows:
    gene = row[1]
    if not isinstance(gene, str) or gene in genes_seen:
        raise ValueError("Missing or duplicate RNA-seq V1 gene ID")
    genes_seen.add(gene)
    values = row[2:]
    if len(values) != 84:
        raise ValueError("RNA-seq row width differs from 84 samples")
    missing_values += sum(value is None for value in values)
    if any(not isinstance(value, (int, float)) or isinstance(value, bool) for value in values):
        raise ValueError(f"Non-numeric RNA-seq expression in {gene}")
    if gene in selected_set:
        selected_rows.append((gene, *values))
if len(genes_seen) != 16606 or missing_values:
    raise ValueError("RNA-seq source row count or completeness failed")
workbook.close()

out_data = ROOT / "data" / "processed" / "external_t007"
out_meta = ROOT / "data" / "metadata" / "external_t007"
out_data.mkdir(parents=True, exist_ok=True)
out_meta.mkdir(parents=True, exist_ok=True)
with (out_data / "PRJNA260535_beta10_log2cpm.tsv").open("w", encoding="utf-8", newline="") as handle:
    writer = csv.writer(handle, delimiter="\t", lineterminator="\n")
    writer.writerow(("Gene", *samples))
    writer.writerows(selected_rows)
with (out_meta / "PRJNA260535_samples.tsv").open("w", encoding="utf-8", newline="") as handle:
    writer = csv.writer(handle, delimiter="\t", lineterminator="\n")
    writer.writerow(("SampleName", "Cultivar", "Brix", "Replicate", "Year", "Tissue"))
    writer.writerows(metadata)
print(f"PRJNA260535 extraction complete: 16606 unique source genes, "
      f"{len(selected_rows)} beta10 genes, 84 skin samples, 0 missing values")
